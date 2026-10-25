"""Leitura das planilhas da ESR (Protheus, template 2027 e projetos_pa) para estruturas simples."""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl
import pandas as pd

from .db import MESES

ABREV_MES = ["jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez"]
# Orçamento de IA para Todos: não liberado em 2026, fica fora da base e pode ser lançado à parte em 2027
ORCAMENTOS_FORA_DA_BASE = {"7016"}


def _txt(serie: pd.Series) -> pd.Series:
    return serie.where(serie.notna(), "").astype(str).str.strip().str.replace(r"\.0$", "", regex=True)


def load_protheus(path: str | Path) -> pd.DataFrame:
    """Base do Protheus com chaves normalizadas (vazio vira texto vazio) e Nº do Projeto recuperado."""
    raw = pd.read_excel(path, sheet_name="BASE ESR", header=None, dtype=object)
    hdr = next(i for i in range(10) if "Tipo de Lançamento" in [str(v) for v in raw.iloc[i]])
    df = raw.iloc[hdr + 1:].copy()
    df.columns = [str(c) for c in raw.iloc[hdr]]
    df = df.loc[:, [c for c in df.columns if c != "nan"]].reset_index(drop=True)
    df["Valor"] = pd.to_numeric(df["Valor"], errors="coerce").fillna(0.0)
    df["Mês"] = df["Mês"].astype(str).str.strip().str.lower().str[:3]
    for c in ["Nº do Orçamento", "Nº do Projeto", "Conta Contábil", "Descrição do Projeto", "Origem de Recurso",
              "Gerente do Projeto", "EDT", "Tarefa"]:
        df[c] = _txt(df[c])
    # orçamentos novos vêm sem Nº do Projeto: tenta recuperar do início da descrição ("2882 - ...")
    da_desc = df["Descrição do Projeto"].str.extract(r"^\s*0*(\d+)\s*-", expand=False).fillna("")
    df["Nº do Projeto"] = df["Nº do Projeto"].where(df["Nº do Projeto"] != "", da_desc)
    return df


def load_contas_template(path: str | Path) -> list[tuple[str, str, str]]:
    """(código, descrição, grupo de despesa) na lista oficial do template."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Conta Contábil"]
    out = []
    for r in ws.iter_rows(min_row=2, max_col=3, values_only=True):
        if r[0] and r[1] and r[2]:
            out.append((str(r[0]).strip(), str(r[1]).strip(), str(r[2]).strip()))
    wb.close()
    return out


def load_listas_template(path: str | Path) -> dict[str, list[str]]:
    """Listas que alimentam as validações do template: diretorias, gerentes e origens."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Diretoria, Gerente, Origem"]
    dir_, ger, ori = [], [], []
    for r in ws.iter_rows(min_row=2, max_col=9, values_only=True):
        if r[1] and r[1] not in dir_:
            dir_.append(str(r[1]))
        if r[3] and r[3] not in ger:
            ger.append(str(r[3]))
        if len(r) > 8 and r[8] and r[8] not in ori:
            ori.append(str(r[8]))
    wb.close()
    return {"diretorias": dir_, "gerentes": ger, "origens": ori}


def load_projetos_pa(path: str | Path) -> dict[str, dict]:
    """Nome oficial (DSS-...) e valor original por Nº do Orçamento (coluna Ação) do projetos_pa."""
    raw = pd.read_excel(path, sheet_name="consolidado_as-is", header=None, dtype=object)
    hdr = next(i for i in range(10) if "ID Projeto" in [str(v) for v in raw.iloc[i]])
    df = raw.iloc[hdr + 1:].copy()
    df.columns = [str(c) for c in raw.iloc[hdr]]
    out: dict[str, dict] = {}
    for _, r in df.iterrows():
        acao = str(r.get("Ação", "")).strip().replace(".0", "")
        if not acao or acao in ("nan", "None") or pd.isna(r.get("Projeto")):
            continue
        out[acao] = {"nome": str(r["Projeto"]).strip(),
                     "original": float(pd.to_numeric(r.get("Valor original"), errors="coerce") or 0),
                     "aprovado": float(pd.to_numeric(r.get("Valor Final Aprovado"), errors="coerce") or 0)}
    return out


def nome_oficial(descricao_protheus: str) -> str:
    """Fallback: '1044 - ESR-ASSOCIACAO-GROWTH' vira 'DSS-ESR-ASSOCIACAO-GROWTH'."""
    resto = descricao_protheus.split(" - ", 1)[-1].strip().lstrip("- ").strip()
    return resto if resto.startswith("DSS-") else f"DSS-{resto}"


def build_projetos(df: pd.DataFrame, pa: dict[str, dict]) -> list[dict]:
    orc = df[df["Tipo de Lançamento"] == "Orçado"]
    out = []
    for numero, g in orc.groupby("Nº do Orçamento", sort=True):
        n_proj = next((v for v in g["Nº do Projeto"] if v), "")
        desc = g["Descrição do Projeto"].iloc[0]
        info = pa.get(numero, {})
        out.append({
            "orcamento": numero,
            "n_projeto": n_proj,
            "descricao": info.get("nome") or nome_oficial(desc),
            "origem": g["Origem de Recurso"].iloc[0],
            "gerente": g["Gerente do Projeto"].iloc[0],
            "original_2026": info.get("original", 0.0),
        })
    return out


def build_linhas_base(df: pd.DataFrame, contas: dict[str, str]) -> list[dict]:
    """Orçado 2026 no grão do template: orçamento, EDT, tarefa, conta e 12 meses."""
    orc = df[(df["Tipo de Lançamento"] == "Orçado") & (~df["Nº do Orçamento"].isin(ORCAMENTOS_FORA_DA_BASE))].copy()
    desconhecidas = sorted(set(orc["Conta Contábil"]) - set(contas))
    if desconhecidas:
        raise ValueError(f"Contas do Protheus fora da lista do template: {desconhecidas}")
    chaves = ["Nº do Orçamento", "EDT", "Tarefa", "Conta Contábil"]
    pv = orc.pivot_table(index=chaves, columns="Mês", values="Valor", aggfunc="sum", fill_value=0.0)
    linhas = []
    for (numero, edt, tarefa, conta), r in pv.iterrows():
        meses = {m: round(float(r.get(ab, 0.0)), 2) for m, ab in zip(MESES, ABREV_MES)}
        linhas.append({"orcamento": numero, "edt": edt, "tarefa": tarefa, "conta_codigo": conta, **meses})
    return linhas


def build_execucao(df: pd.DataFrame) -> list[dict]:
    """Realizado e empenhado de 2026 por orçamento, EDT, tarefa e conta (valores positivos)."""
    ex = df[df["Tipo de Lançamento"].isin(["Realizado", "Empenhado"])]
    chaves = ["Nº do Orçamento", "EDT", "Tarefa", "Conta Contábil"]
    pv = ex.pivot_table(index=chaves, columns="Tipo de Lançamento", values="Valor", aggfunc="sum", fill_value=0.0)
    out = []
    for (numero, edt, tarefa, conta), r in pv.iterrows():
        out.append({"orcamento": numero, "edt": edt, "tarefa": tarefa, "conta_codigo": conta,
                    "realizado": round(-float(r.get("Realizado", 0.0)), 2),
                    "empenhado": round(-float(r.get("Empenhado", 0.0)), 2)})
    return out


def sem_acento(s: str) -> str:
    import unicodedata
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def login_de(nome: str) -> str:
    """'Luciana Batista' vira 'luciana'."""
    return re.sub(r"[^a-z0-9]", "", sem_acento(nome.split()[0]).lower())
