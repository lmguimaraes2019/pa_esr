"""Planejamento de viagens: cálculo de custos, carga a partir da planilha e ligação com o orçamento.

Regras (as mesmas da planilha VIAGENS - ORÇAMENTO 2026):
  pernoites   = dias - 1 (pode ser ajustado à mão na viagem)
  passagem    = passagem(ida/volta) x viagens + bagagem x viagens   (a bagagem entra na conta de passagens)
  hospedagem  = pernoites x diária de hospedagem x viagens
  ajuda custo = ajuda de custo por dia x dias x viagens
  transporte  = auxílio transporte x viagens
  seguro      = seguro x viagens (internacionais)
Viagens com menos de 1 dia não geram custo.
"""
from __future__ import annotations

import sqlite3
import unicodedata
from pathlib import Path

from . import config
from .db import MESES, audit

NOMES_MES = ["JANEIRO", "FEVEREIRO", "MARÇO", "ABRIL", "MAIO", "JUNHO", "JULHO", "AGOSTO",
             "SETEMBRO", "OUTUBRO", "NOVEMBRO", "DEZEMBRO"]

# componente do custo -> conta contábil do template
CONTAS = {
    "nac": {"passagem": "51030504", "hospedagem": "51030503", "ajuda_custo": "51030502", "aux_transporte": "51030506"},
    "int": {"passagem": "51030513", "hospedagem": "51030512", "ajuda_custo": "51030511",
            "aux_transporte": "51030515", "seguro": "51030514"},
}
CONTAS_VIAGEM = {c for d in CONTAS.values() for c in d.values()}
COMPONENTES = ["passagem", "hospedagem", "ajuda_custo", "aux_transporte", "seguro"]
CAMPOS_TARIFA = ["passagem", "hospedagem", "bagagem", "aux_transporte", "ajuda_custo", "seguro"]


def calcular(tipo: str, t: sqlite3.Row | dict | None, n: int, dias: int, pernoites: int | None = None,
             bagagem_vezes: float = 1.0) -> dict[str, float]:
    """Custo de uma linha de viagens, por componente. Sem tarifa ou sem dias, tudo zero."""
    zero = {k: 0.0 for k in COMPONENTES} | {"bagagem": 0.0, "total": 0.0, "pernoites": 0}
    if t is None or dias < 1 or n < 1:
        return zero
    pern = max(dias - 1, 0) if pernoites is None else max(int(pernoites), 0)
    bag = float(t["bagagem"]) * n * bagagem_vezes
    out = {
        "passagem": float(t["passagem"]) * n + bag,
        "hospedagem": pern * float(t["hospedagem"]) * n,
        "ajuda_custo": float(t["ajuda_custo"]) * dias * n,
        "aux_transporte": float(t["aux_transporte"]) * n,
        "seguro": float(t["seguro"]) * n if tipo == "int" else 0.0,
        "bagagem": bag,
        "pernoites": pern,
    }
    out["total"] = sum(out[k] for k in COMPONENTES)
    return {k: (round(v, 2) if k != "pernoites" else v) for k, v in out.items()}


def bagagem_vezes(c: sqlite3.Connection) -> float:
    r = c.execute("SELECT valor FROM premissas WHERE chave='viagens_bagagem_vezes'").fetchone()
    return float(r[0]) if r else 1.0


def tarifas_map(c: sqlite3.Connection) -> dict[tuple[str, str, str], sqlite3.Row]:
    return {(r["tipo"], r["origem"], r["destino"]): r for r in c.execute("SELECT * FROM viagem_tarifas")}


# ---------------------------------------------------------------- carga da planilha

def _num(v) -> float:
    try:
        return float(v) if v not in (None, "") else 0.0
    except (TypeError, ValueError):
        return 0.0


def _s(v) -> str:
    return "" if v is None else str(v).strip()


def _mes(nome: str) -> int | None:
    n = unicodedata.normalize("NFC", _s(nome).upper())
    return NOMES_MES.index(n) + 1 if n in NOMES_MES else None


def _pern(valor, dias) -> int | None:
    """Pernoites digitados à mão na planilha (diferentes de dias - 1); caso contrário, vazio."""
    p, d = int(_num(valor)), int(_num(dias))
    return p if p != max(d - 1, 0) else None


def ler_planilha(caminho: Path) -> dict:
    import warnings
    import openpyxl
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(caminho, data_only=True)
    ap = wb["APOIO"]
    tarifas = []
    for r in ap.iter_rows(min_row=5, max_row=60, values_only=True):
        # nacionais: B..H (índices 1..7); internacionais: P..Y (índices 15..24)
        if _s(r[1]) and _s(r[2]) and r[3] is not None:
            tarifas.append({"tipo": "nac", "origem": _s(r[1]), "destino": _s(r[2]), "passagem": _num(r[3]),
                            "hospedagem": _num(r[4]), "bagagem": _num(r[5]), "aux_transporte": _num(r[6]),
                            "ajuda_custo": _num(r[7]), "seguro": 0.0})
        if len(r) > 24 and _s(r[15]) and _s(r[16]) and r[17] is not None:
            tarifas.append({"tipo": "int", "origem": _s(r[15]), "destino": _s(r[16]), "passagem": _num(r[17]),
                            "hospedagem": _num(r[18]), "bagagem": _num(r[19]), "aux_transporte": _num(r[20]),
                            "ajuda_custo": _num(r[21]), "seguro": _num(r[24])})
    viagens = []
    for r in wb["VIAGENS NACIONAIS"].iter_rows(min_row=7, max_row=200, values_only=True):
        m = _mes(r[1])
        if m and _s(r[2]) and _s(r[3]):
            viagens.append({"tipo": "nac", "mes": m, "origem": _s(r[2]), "destino": _s(r[3]),
                            "n_viagens": int(_num(r[4]) or 1), "dias": int(_num(r[5])), "pernoites": _pern(r[6], r[5]), "objetivo": _s(r[13]),
                            "area": _s(r[14]), "orcamento": _s(int(r[15])) if r[15] not in (None, "") else ""})
    for r in wb["VIAGENS INTERNACIONAIS"].iter_rows(min_row=7, max_row=200, values_only=True):
        m = _mes(r[2])
        if m and _s(r[3]) and _s(r[4]):
            viagens.append({"tipo": "int", "mes": m, "origem": _s(r[3]), "destino": _s(r[4]),
                            "n_viagens": int(_num(r[5]) or 1), "dias": int(_num(r[6])), "pernoites": _pern(r[7], r[6]), "objetivo": _s(r[15]),
                            "area": "", "orcamento": ""})
    return {"tarifas": tarifas, "viagens": viagens}


def tarefa_padrao(c: sqlite3.Connection, orcamento: str) -> tuple[str, str] | None:
    """Tarefa de viagens do projeto no orçamento aprovado, se for uma só."""
    rows = c.execute(
        "SELECT edt, tarefa, SUM(" + "+".join(MESES) + ") t FROM linhas WHERE ciclo=? AND orcamento=? "
        "AND UPPER(tarefa) LIKE '%VIAGE%' GROUP BY edt, tarefa", (config.CICLO_BASE, orcamento)).fetchall()
    com_valor = [r for r in rows if r["t"] > 1]
    escolha = com_valor or rows
    return (escolha[0]["edt"], escolha[0]["tarefa"]) if len(escolha) == 1 else None


def importar(c: sqlite3.Connection, caminho: Path, reset: bool = False) -> dict:
    if c.execute("SELECT COUNT(*) FROM viagens").fetchone()[0]:
        if not reset:
            raise SystemExit("Já existem viagens no banco. Use --reset para recarregar.")
        c.execute("DELETE FROM viagens")
        c.execute("DELETE FROM viagem_tarifas")
        c.execute("UPDATE linhas SET fonte='manual'")
    dados = ler_planilha(caminho)
    for t in dados["tarifas"]:
        vals = [t[k] for k in CAMPOS_TARIFA]
        c.execute("INSERT INTO viagem_tarifas (tipo, origem, destino, passagem, hospedagem, bagagem, aux_transporte,"
                  " ajuda_custo, seguro, ref_passagem, ref_hospedagem, ref_bagagem, ref_aux_transporte,"
                  " ref_ajuda_custo, ref_seguro) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                  (t["tipo"], t["origem"], t["destino"], *vals, *vals))
    projetos = {r[0] for r in c.execute("SELECT orcamento FROM projetos")}
    padrao: dict[str, tuple[str, str] | None] = {}
    sem_projeto = 0
    for v in dados["viagens"]:
        o = v["orcamento"] if v["orcamento"] in projetos else None
        if v["orcamento"] and not o:
            sem_projeto += 1
        if o and o not in padrao:
            padrao[o] = tarefa_padrao(c, o)
        edt, tarefa = (padrao.get(o) or (None, None)) if o else (None, None)
        cur = c.execute("INSERT INTO viagens (ciclo, tipo, mes, origem, destino, n_viagens, dias, pernoites, objetivo, area,"
                        " orcamento, edt, tarefa, atualizado_por) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                        (config.CICLO_BASE, v["tipo"], v["mes"], v["origem"], v["destino"], v["n_viagens"], v["dias"],
                         v["pernoites"], v["objetivo"], v["area"], o, edt, tarefa, "carga inicial"))
        c.execute("INSERT INTO viagens (ciclo, tipo, mes, origem, destino, n_viagens, dias, pernoites, objetivo, area,"
                  " orcamento, edt, tarefa, base_id, atualizado_por) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                  (config.CICLO_PLANO, v["tipo"], v["mes"], v["origem"], v["destino"], v["n_viagens"], v["dias"],
                   v["pernoites"], v["objetivo"], v["area"], o, edt, tarefa, cur.lastrowid, "carga inicial"))
    return {"tarifas": len(dados["tarifas"]), "viagens": len(dados["viagens"]), "projeto_desconhecido": sem_projeto}


# ---------------------------------------------------------------- plano -> orçamento

def plano_por_linha(c: sqlite3.Connection, ciclo: int = config.CICLO_PLANO):
    """Soma o plano de viagens por (orçamento, edt, tarefa, conta) e mês. Devolve (plano, pendências)."""
    tar = tarifas_map(c)
    bvezes = bagagem_vezes(c)
    plano: dict[tuple[str, str, str, str], list[float]] = {}
    pend: list[str] = []
    for v in c.execute("SELECT * FROM viagens WHERE ciclo=? ORDER BY id", (ciclo,)):
        t = tar.get((v["tipo"], v["origem"], v["destino"]))
        custo = calcular(v["tipo"], t, v["n_viagens"], v["dias"], v["pernoites"], bvezes)
        if t is None and v["dias"] >= 1:
            pend.append(f"Viagem {v['id']} ({v['origem']} para {v['destino']}) sem tarifa cadastrada.")
            continue
        if custo["total"] <= 0:
            continue
        if not v["orcamento"] or not v["edt"] or not v["tarefa"]:
            pend.append(f"Viagem {v['id']} ({v['objetivo'][:50] or v['destino']}) sem projeto ou tarefa do orçamento.")
            continue
        for comp, conta in CONTAS[v["tipo"]].items():
            val = custo[comp]
            if val:
                linha = plano.setdefault((v["orcamento"], v["edt"], v["tarefa"], conta), [0.0] * 12)
                linha[v["mes"] - 1] = round(linha[v["mes"] - 1] + val, 2)
    return plano, pend


def aplicar_ao_orcamento(c: sqlite3.Connection, usuario: str, simular: bool = True) -> dict:
    plano, pend = plano_por_linha(c)
    tarefas = {(r["orcamento"], r["edt"], r["tarefa"]) for r in c.execute(
        "SELECT DISTINCT orcamento, edt, tarefa FROM viagens WHERE orcamento IS NOT NULL AND edt IS NOT NULL "
        "AND tarefa IS NOT NULL")}
    tarefas |= {k[:3] for k in plano}
    tarefas |= {(r["orcamento"], r["edt"], r["tarefa"]) for r in c.execute(
        "SELECT DISTINCT orcamento, edt, tarefa FROM linhas WHERE ciclo=? AND fonte='viagens'", (config.CICLO_PLANO,))}
    erros = list(pend)
    existentes = {(r["orcamento"], r["edt"], r["tarefa"]) for r in c.execute(
        "SELECT DISTINCT orcamento, edt, tarefa FROM linhas")}
    for t in sorted(tarefas):
        if t not in existentes:
            erros.append(f"A tarefa {t[2]} do orçamento {t[0]} não existe no projeto.")
    if erros:
        return {"bloqueado": True, "erros": erros, "itens": [], "total_antes": 0, "total_depois": 0}

    itens, total_antes, total_depois = [], 0.0, 0.0
    soma = "+".join(MESES)
    for (o, edt, tarefa) in sorted(tarefas):
        linhas = c.execute(
            f"SELECT l.*, ({soma}) AS tot, (SELECT COALESCE(SUM({soma}),0) FROM linhas b WHERE b.id = l.base_id) AS ref "
            "FROM linhas l WHERE l.ciclo=? AND l.orcamento=? AND l.edt=? AND l.tarefa=? AND l.fonte NOT IN ('eventos','headcount') AND l.conta_codigo IN (%s)"
            % ",".join("?" for _ in CONTAS_VIAGEM), (config.CICLO_PLANO, o, edt, tarefa, *sorted(CONTAS_VIAGEM))).fetchall()
        por_conta = {r["conta_codigo"]: r for r in linhas}
        contas = sorted(set(por_conta) | {k[3] for k in plano if k[:3] == (o, edt, tarefa)})
        for conta in contas:
            novo = plano.get((o, edt, tarefa, conta), [0.0] * 12)
            atual = por_conta.get(conta)
            antes = round(atual["tot"], 2) if atual else 0.0
            depois = round(sum(novo), 2)
            total_antes += antes
            total_depois += depois
            descricao = c.execute("SELECT descricao FROM contas WHERE codigo=?", (conta,)).fetchone()[0]
            itens.append({"orcamento": o, "edt": edt, "tarefa": tarefa, "conta": conta, "conta_desc": descricao,
                          "ref_2026": round(atual["ref"], 2) if atual else 0.0, "antes": antes, "depois": depois,
                          "manual": bool(atual and atual["fonte"] != "viagens" and antes > 0.005)})
            if simular:
                continue
            if atual:
                if depois == 0 and atual["base_id"] is None:
                    c.execute("DELETE FROM linhas WHERE id=?", (atual["id"],))
                else:
                    c.execute(f"UPDATE linhas SET {', '.join(m + '=?' for m in MESES)}, fonte='viagens', editada=0, "
                              "atualizado_por=?, atualizado_em=datetime('now') WHERE id=?", (*novo, usuario, atual["id"]))
            elif depois > 0:
                c.execute(f"INSERT INTO linhas (ciclo, orcamento, edt, tarefa, conta_codigo, {', '.join(MESES)}, fonte, "
                          "atualizado_por) VALUES (?,?,?,?,?,%s,'viagens',?)" % ",".join("?" for _ in MESES),
                          (config.CICLO_PLANO, o, edt, tarefa, conta, *novo, usuario))
    if not simular:
        audit(c, usuario, "aplicou plano de viagens ao orçamento",
              detalhe=f"{len(itens)} linhas; de {round(total_antes, 2)} para {round(total_depois, 2)}")
    return {"bloqueado": False, "erros": [], "itens": itens, "total_antes": round(total_antes, 2),
            "total_depois": round(total_depois, 2)}
