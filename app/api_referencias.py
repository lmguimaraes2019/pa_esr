"""Administração: importação das tabelas de referência (tarifas de viagens, catálogo de equipamentos e tabela salarial).

Cada importação lê uma planilha .xlsx ou .csv, mostra o que muda (conferência) e só grava quando o administrador confirma.
O arquivo modelo de cada tabela já traz os valores vigentes: basta baixar, editar e importar de volta.
"""
from __future__ import annotations

import csv
import io
import re
import sqlite3
import unicodedata
from datetime import datetime

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import Response

from . import equipamentos, pessoal, sync, viagens as vg
from .api import _conn, exige_admin
from .db import audit, tx

router = APIRouter(prefix="/api/referencias")
LIMITE = 10 * 1024 * 1024
AMOSTRA = 300

TABELAS = {
    "viagens": {"nome": "Tarifas de viagens", "colunas": ["tipo", "origem", "destino", *vg.CAMPOS_TARIFA]},
    "equipamentos": {"nome": "Catálogo de equipamentos", "colunas": ["codigo", "tipo", "nome", "perfil", "valor"]},
    "salarios": {"nome": "Tabela salarial", "colunas": ["trajetoria", "nivel", "grupo", "inicio", "fim"]},
}
ROTULOS = {
    "tipo": "Tipo", "origem": "Origem", "destino": "Destino", "passagem": "Passagem", "hospedagem": "Hospedagem", "bagagem": "Bagagem",
    "aux_transporte": "Aux. transporte", "ajuda_custo": "Ajuda de custo", "seguro": "Seguro", "codigo": "Código", "nome": "Equipamento",
    "perfil": "Perfil", "valor": "Valor", "trajetoria": "Trajetória", "nivel": "Nível", "grupo": "Grupo", "inicio": "Início da faixa (80%)",
    "fim": "Fim da faixa (120%)",
}
ALIAS = {
    "tipo": ["tipo", "categoria"], "origem": ["origem"], "destino": ["destino"],
    "passagem": ["passagem"], "hospedagem": ["hospedagem"], "bagagem": ["bagagem"],
    "aux_transporte": ["aux_transporte", "aux._transporte", "auxilio_transporte", "aux_transp"], "ajuda_custo": ["ajuda_custo", "ajuda_de_custo"], "seguro": ["seguro"],
    "codigo": ["codigo", "nivel_modelo", "cod"], "nome": ["nome", "equipamento", "modelo", "equipamento_modelo", "item", "descricao"],
    "perfil": ["perfil", "observacao", "perfil_observacao"], "valor": ["valor", "valor_r", "valor_(r$)", "preco"],
    "trajetoria": ["trajetoria"], "nivel": ["nivel"], "grupo": ["grupo"],
    "inicio": ["inicio", "inicio_da_faixa", "inicio_da_faixa_(80%)", "faixa_inicio", "minimo"], "fim": ["fim", "fim_da_faixa", "fim_da_faixa_(120%)", "faixa_fim", "maximo"],
    "referencia": ["referencia", "valor_100%", "ponto_medio", "salario_referencia", "100%"],
}


def _sem_acento(t: str) -> str:
    return "".join(ch for ch in unicodedata.normalize("NFD", str(t)) if unicodedata.category(ch) != "Mn")


def _chave(t) -> str:
    return re.sub(r"[\s]+", "_", _sem_acento(str(t or "")).strip().lower())


def _numero(v) -> float | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = re.sub(r"[^\d,.\-]", "", str(v))
    if not t or t in "-.,":
        return None
    if "," in t and "." in t:
        t = t.replace(".", "").replace(",", ".")
    elif "," in t:
        t = t.replace(",", ".")
    try:
        return float(t)
    except ValueError:
        return None


def _ler(nome: str, dados: bytes) -> list[dict]:
    """Linhas do arquivo como dicionários com os cabeçalhos normalizados."""
    if len(dados) > LIMITE:
        raise HTTPException(413, "Arquivo grande demais (limite de 10 MB).")
    if dados[:2] == b"PK":
        import openpyxl
        import warnings
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb = openpyxl.load_workbook(io.BytesIO(dados), read_only=True, data_only=True)
        linhas = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    elif nome.lower().endswith((".csv", ".txt")):
        texto = dados.decode("utf-8-sig", errors="replace") if b"\xef\xbf\xbd" not in dados else dados.decode("latin-1")
        try:
            dialeto = csv.Sniffer().sniff(texto[:4096], delimiters=";,\t")
        except csv.Error:
            dialeto = csv.excel
        linhas = list(csv.reader(io.StringIO(texto), dialeto))
    else:
        raise HTTPException(422, "Envie uma planilha .xlsx ou um arquivo .csv.")
    linhas = [l for l in linhas if any(str(x).strip() for x in l if x is not None)]
    if len(linhas) < 2:
        raise HTTPException(422, "O arquivo não tem linhas de dados.")
    cab = [_chave(x) for x in linhas[0]]
    return [{cab[i]: (v if not isinstance(v, str) else v.strip()) for i, v in enumerate(l) if i < len(cab) and cab[i]} for l in linhas[1:]]


def _pega(linha: dict, campo: str):
    for a in ALIAS[campo]:
        if a in linha and linha[a] not in (None, ""):
            return linha[a]
    return None


def _exige_colunas(linhas: list[dict], campos: list[str]) -> None:
    presentes = set().union(*(l.keys() for l in linhas))
    faltam = [ROTULOS[c] for c in campos if not any(a in presentes for a in ALIAS[c])]
    if faltam:
        raise HTTPException(422, "Colunas não encontradas na primeira linha do arquivo: " + ", ".join(faltam) + ". Baixe o modelo para ver os cabeçalhos.")


# ---------------------------------------------------------------- leitura e conferência de cada tabela

def _tipo_viagem(v) -> str | None:
    t = _chave(v)
    return "nac" if t in ("nac", "nacional") else "int" if t in ("int", "internacional") else None


def _preparar_viagens(c: sqlite3.Connection, linhas: list[dict]) -> dict:
    _exige_colunas(linhas, ["tipo", "origem", "destino", "passagem"])
    atuais = {(r["tipo"], r["origem"].upper(), r["destino"].upper()): r for r in c.execute("SELECT * FROM viagem_tarifas")}
    novos, erros, vistos = [], [], set()
    for n, l in enumerate(linhas, start=2):
        tipo, o, d = _tipo_viagem(_pega(l, "tipo")), str(_pega(l, "origem") or "").strip(), str(_pega(l, "destino") or "").strip()
        if not (tipo and o and d):
            erros.append(f"Linha {n}: informe tipo (nacional ou internacional), origem e destino.")
            continue
        chave = (tipo, o.upper(), d.upper())
        if chave in vistos:
            erros.append(f"Linha {n}: tarifa repetida ({o} para {d}).")
            continue
        vistos.add(chave)
        vals = {}
        for k in vg.CAMPOS_TARIFA:
            x = _pega(l, k)
            v = _numero(x) if x is not None else 0.0
            if v is None or v < 0 or v > 1e7:
                erros.append(f"Linha {n}: valor inválido em {ROTULOS[k]}.")
                v = 0.0
            vals[k] = round(v, 2)
        novos.append({"chave": chave, "tipo": tipo, "origem": o, "destino": d, **vals})
    mudancas, iguais, add = [], 0, 0
    for x in novos:
        a = atuais.get(x["chave"])
        rot = f"{'Nacional' if x['tipo'] == 'nac' else 'Internacional'}: {x['origem']} para {x['destino']}"
        if not a:
            add += 1
            mudancas.append({"item": rot, "situacao": "novo", "antes": "", "depois": ", ".join(f"{ROTULOS[k]} {x[k]:.2f}" for k in vg.CAMPOS_TARIFA if x[k])})
        else:
            dif = [k for k in vg.CAMPOS_TARIFA if abs(a[k] - x[k]) >= 0.005]
            if dif:
                mudancas.append({"item": rot, "situacao": "alterado", "antes": ", ".join(f"{ROTULOS[k]} {a[k]:.2f}" for k in dif),
                                 "depois": ", ".join(f"{ROTULOS[k]} {x[k]:.2f}" for k in dif)})
            else:
                iguais += 1
    ausentes = [f"{'Nacional' if k[0] == 'nac' else 'Internacional'}: {r['origem']} para {r['destino']}" for k, r in atuais.items() if k not in vistos]
    return {"lidos": len(novos), "novos": add, "alterados": len(mudancas) - add, "iguais": iguais, "ausentes": ausentes, "mudancas": mudancas,
            "erros": erros, "aviso": "As tarifas que não constam no arquivo continuam como estão." if ausentes else "", "_dados": novos, "_atuais": atuais}


def _aplicar_viagens(c: sqlite3.Connection, p: dict) -> None:
    for x in p["_dados"]:
        a = p["_atuais"].get(x["chave"])
        if a:
            c.execute(f"UPDATE viagem_tarifas SET {', '.join(k + '=?' for k in vg.CAMPOS_TARIFA)} WHERE id=?", (*[x[k] for k in vg.CAMPOS_TARIFA], a["id"]))
        else:
            vals = [x[k] for k in vg.CAMPOS_TARIFA]
            c.execute("INSERT INTO viagem_tarifas (tipo, origem, destino, passagem, hospedagem, bagagem, aux_transporte, ajuda_custo, seguro, "
                      "ref_passagem, ref_hospedagem, ref_bagagem, ref_aux_transporte, ref_ajuda_custo, ref_seguro) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                      (x["tipo"], x["origem"], x["destino"], *vals, *vals))


_TIPOS_EQ = {"computador": "computador", "notebook": "computador", "desktop": "computador", "monitor": "monitor", "monitores": "monitor",
             "periferico": "periferico", "acessorio": "periferico", "dock": "periferico", "periferico_acessorio": "periferico"}


def _preparar_equipamentos(c: sqlite3.Connection, linhas: list[dict]) -> dict:
    _exige_colunas(linhas, ["codigo", "tipo", "nome", "valor"])
    atuais = {i["codigo"]: i for i in equipamentos.catalogo()}
    novos, erros, vistos = [], [], set()
    for n, l in enumerate(linhas, start=2):
        cod, nome = str(_pega(l, "codigo") or "").strip(), str(_pega(l, "nome") or "").strip()
        tipo = _TIPOS_EQ.get(_chave(_pega(l, "tipo")))
        valor = _numero(_pega(l, "valor"))
        if not (cod and nome):
            erros.append(f"Linha {n}: informe código e equipamento.")
        elif not tipo:
            erros.append(f"Linha {n} ({cod}): tipo deve ser computador, monitor ou periférico.")
        elif valor is None or valor < 0 or valor > 1e7:
            erros.append(f"Linha {n} ({cod}): valor inválido.")
        elif cod in vistos:
            erros.append(f"Linha {n}: código {cod} repetido.")
        else:
            vistos.add(cod)
            novos.append({"codigo": cod, "tipo": tipo, "nome": nome, "perfil": str(_pega(l, "perfil") or "").strip(), "valor": round(valor, 2)})
    mudancas, iguais = [], 0
    for x in novos:
        a = atuais.get(x["codigo"])
        if not a:
            mudancas.append({"item": f"{x['codigo']} | {x['nome']}", "situacao": "novo", "antes": "", "depois": f"{x['tipo']}, R$ {x['valor']:.2f}"})
        else:
            dif = [k for k in ("tipo", "nome", "perfil", "valor") if a.get(k) != x[k]]
            if dif:
                mudancas.append({"item": f"{x['codigo']} | {x['nome']}", "situacao": "alterado",
                                 "antes": "; ".join(f"{k} {a.get(k)}" for k in dif), "depois": "; ".join(f"{k} {x[k]}" for k in dif)})
            else:
                iguais += 1
    ausentes = [f"{k} | {a['nome']}" for k, a in atuais.items() if k not in vistos]
    usados = {r[0] for r in c.execute("SELECT computador FROM headcount UNION SELECT monitor FROM headcount UNION SELECT periferico FROM headcount") if r[0]}
    for cod in sorted(usados - vistos):
        erros.append(f"O equipamento {cod} está em posições de headcount e não consta no arquivo. Mantenha-o no catálogo ou ajuste as posições antes.")
    return {"lidos": len(novos), "novos": sum(1 for m in mudancas if m["situacao"] == "novo"), "alterados": sum(1 for m in mudancas if m["situacao"] == "alterado"),
            "iguais": iguais, "ausentes": ausentes, "mudancas": mudancas, "erros": erros,
            "aviso": "Os equipamentos que não constam no arquivo saem do catálogo." if ausentes else "", "_dados": novos}


def _norm_traj(t) -> str:
    return re.sub(r"\s+", " ", _sem_acento(str(t or "")).strip().upper())


def _preparar_salarios(c: sqlite3.Connection, linhas: list[dict]) -> dict:
    _exige_colunas(linhas, ["trajetoria", "nivel", "grupo"])
    presentes = set().union(*(l.keys() for l in linhas))
    tem_faixa = any(a in presentes for a in ALIAS["inicio"]) and any(a in presentes for a in ALIAS["fim"])
    if not tem_faixa and not any(a in presentes for a in ALIAS["referencia"]):
        raise HTTPException(422, "Informe o início e o fim da faixa, ou uma coluna de salário de referência (100%).")
    atuais = {(f["trajetoria"], f["nivel"], f["grupo"]): f for f in pessoal.tabela_salarial()}
    novos, erros, vistos = [], [], set()
    for n, l in enumerate(linhas, start=2):
        t, nv, g = _norm_traj(_pega(l, "trajetoria")), _norm_traj(_pega(l, "nivel")), _norm_traj(_pega(l, "grupo"))
        ini, fim, ref = _numero(_pega(l, "inicio")), _numero(_pega(l, "fim")), _numero(_pega(l, "referencia"))
        if ini is None and fim is None and ref is not None:
            ini, fim = round(ref * 0.8, 2), round(ref * 1.2, 2)
        if not (t and nv and g):
            erros.append(f"Linha {n}: informe trajetória, nível e grupo.")
        elif ini is None or fim is None or ini <= 0 or fim < ini or fim > 1e7:
            erros.append(f"Linha {n} ({t} / {nv} / {g}): faixa inválida (o fim deve ser maior ou igual ao início).")
        elif (t, nv, g) in vistos:
            erros.append(f"Linha {n}: {t} / {nv} / {g} repetido.")
        else:
            vistos.add((t, nv, g))
            novos.append({"trajetoria": t, "nivel": nv, "grupo": g, "inicio": round(ini, 2), "fim": round(fim, 2)})
    mudancas, iguais = [], 0
    for x in novos:
        a = atuais.get((x["trajetoria"], x["nivel"], x["grupo"]))
        rot = f"{x['trajetoria']} / {x['nivel']} / {x['grupo']}"
        if not a:
            mudancas.append({"item": rot, "situacao": "novo", "antes": "", "depois": f"{x['inicio']:.2f} a {x['fim']:.2f}"})
        elif abs(a["inicio"] - x["inicio"]) >= 0.005 or abs(a["fim"] - x["fim"]) >= 0.005:
            mudancas.append({"item": rot, "situacao": "alterado", "antes": f"{a['inicio']:.2f} a {a['fim']:.2f}", "depois": f"{x['inicio']:.2f} a {x['fim']:.2f}"})
        else:
            iguais += 1
    ausentes = [" / ".join(k) for k in atuais if k not in vistos]
    for r in c.execute("SELECT DISTINCT trajetoria, nivel, grupo FROM headcount"):
        if (r[0], r[1], r[2]) not in vistos:
            erros.append(f"A faixa {r[0]} / {r[1]} / {r[2]} está em posições de headcount e não consta no arquivo. Mantenha-a na tabela ou ajuste as posições antes.")
    return {"lidos": len(novos), "novos": sum(1 for m in mudancas if m["situacao"] == "novo"), "alterados": sum(1 for m in mudancas if m["situacao"] == "alterado"),
            "iguais": iguais, "ausentes": ausentes, "mudancas": mudancas, "erros": erros,
            "aviso": "As faixas que não constam no arquivo saem da tabela." if ausentes else "", "_dados": novos}


PREPARAR = {"viagens": _preparar_viagens, "equipamentos": _preparar_equipamentos, "salarios": _preparar_salarios}


# ---------------------------------------------------------------- rotas

def _tabela(tipo: str) -> dict:
    if tipo not in TABELAS:
        raise HTTPException(404, "Tabela não encontrada.")
    return TABELAS[tipo]


@router.get("")
def situacao(_: dict = Depends(exige_admin), c: sqlite3.Connection = Depends(_conn)):
    eq = equipamentos.carregar()
    ti = pessoal.tabela_info()
    return [
        {"tipo": "viagens", "nome": TABELAS["viagens"]["nome"], "registros": c.execute("SELECT COUNT(*) FROM viagem_tarifas").fetchone()[0],
         "origem": "banco de dados", "detalhe": "Passagem, hospedagem, bagagem, auxílio de transporte, ajuda de custo e seguro por rota.",
         "colunas": [ROTULOS[k] for k in TABELAS["viagens"]["colunas"]]},
        {"tipo": "equipamentos", "nome": TABELAS["equipamentos"]["nome"], "registros": len(eq["itens"]), "origem": eq.get("origem", "importado"),
         "detalhe": f"Catálogo de Hardware Homologado, atualizado em {eq['atualizado_em']}.", "colunas": [ROTULOS[k] for k in TABELAS["equipamentos"]["colunas"]]},
        {"tipo": "salarios", "nome": TABELAS["salarios"]["nome"], "registros": len(ti["faixas"]), "origem": ti["origem"],
         "detalhe": f"{ti['nome']}: faixas por trajetória, nível e grupo (80% a 120% da referência).", "colunas": [ROTULOS[k] for k in TABELAS["salarios"]["colunas"]]},
    ]


@router.get("/{tipo}/modelo")
def modelo(tipo: str, _: dict = Depends(exige_admin), c: sqlite3.Connection = Depends(_conn)):
    """Planilha com os valores vigentes, para baixar, editar e importar de volta."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    cols = _tabela(tipo)["colunas"]
    if tipo == "viagens":
        linhas = [[("nacional" if r["tipo"] == "nac" else "internacional") if k == "tipo" else r[k] for k in cols]
                  for r in c.execute("SELECT * FROM viagem_tarifas ORDER BY tipo, origem, destino")]
    elif tipo == "equipamentos":
        linhas = [[i[k] for k in cols] for i in equipamentos.catalogo()]
    else:
        linhas = [[f[k] for k in cols] for f in pessoal.tabela_salarial()]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = TABELAS[tipo]["nome"][:30]
    ws.append([ROTULOS[k] for k in cols])
    for cel in ws[1]:
        cel.font, cel.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="003366")
    for l in linhas:
        ws.append(l)
    for i, k in enumerate(cols, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = 60 if k == "nome" else 26
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="modelo_{tipo}.xlsx"'})


@router.post("/{tipo}/importar")
async def importar(tipo: str, simular: bool = True, arquivo: UploadFile = File(...), u: dict = Depends(exige_admin)):
    _tabela(tipo)
    linhas = _ler(arquivo.filename or "", await arquivo.read())
    with tx() as c:
        p = PREPARAR[tipo](c, linhas)
        limpo = {k: v for k, v in p.items() if not k.startswith("_")}
        limpo.update(simulacao=simular, tipo=tipo, mudancas=p["mudancas"][:AMOSTRA], mais_mudancas=max(0, len(p["mudancas"]) - AMOSTRA),
                     ausentes=p["ausentes"][:50], mais_ausentes=max(0, len(p["ausentes"]) - 50), erros=p["erros"][:50], mais_erros=max(0, len(p["erros"]) - 50))
        if simular:
            return limpo
        if p["erros"]:
            raise HTTPException(422, "Há erros no arquivo. Corrija e confira de novo antes de importar:\n" + "\n".join(p["erros"][:10]))
        hoje = datetime.now().strftime("%d/%m/%Y")
        if tipo == "viagens":
            _aplicar_viagens(c, p)
        elif tipo == "equipamentos":
            equipamentos.gravar(p["_dados"], hoje, f"importado de {arquivo.filename}")
        else:
            pessoal.gravar_tabela(p["_dados"], f"Tabela Salarial importada de {arquivo.filename} em {hoje}")
        audit(c, u["login"], f"importou {TABELAS[tipo]['nome'].lower()}",
              detalhe=f"{arquivo.filename}: {limpo['novos']} novos, {limpo['alterados']} alterados, {len(p['ausentes'])} ausentes do arquivo")
        sync.sincronizar_todos(c, u["login"])      # custos de viagens, eventos e headcount passam a usar os valores novos
        return limpo
